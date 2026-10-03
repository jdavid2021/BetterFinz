import hashlib, hmac, io, logging, re, time, uuid
from datetime import date, timedelta
from decimal import Decimal
from urllib.parse import urlparse
from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.middleware.trustedhost import TrustedHostMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from prometheus_client import CONTENT_TYPE_LATEST, generate_latest
from sqlalchemy import case, func, or_, select
from sqlalchemy.orm import Session
from app.auth import COOKIE, current_user, household_id
from app.account_merge_service import AccountMergeError, merge_accounts
from app.account_service import (
    archive_account,
    create_manual_account,
    update_manual_account,
    update_manual_connection,
)
from app.account_history_service import account_history
from app.account_onboarding_service import preview_account_statement
from app.balance_service import BalanceView, available_balance_view
from app.bill_service import bill_candidates, create_bill, ignore_bill_candidate, match_open_bill_payments, payment_match_candidates, transaction_payment_contexts, update_bill
from app.biller_directory import match_biller, public_biller, search_billers
from app.institution_directory import get_institution, search_institutions
from app.config import settings
from app.db import get_db
from app.health import probe_dependencies, refresh_health_metrics
from app.metrics import API_ERRORS, API_LATENCY, API_REQUESTS
from app.observability import request_id_context
from app.categorization_service import learn_from_transaction
from app.accounting_service import (
    public_categories,
    AccountingError,
    category_for,
    ensure_household_accounting,
    update_transaction_categories as post_transaction_categories,
)
from app.debt_service import (
    create_statement,
    generate_strategy,
    import_liability_pdf,
    latest_statements,
    saved_strategy,
    schedule_liability,
)
from app.statement_service import import_statement
from app.monthly_plan_service import monthly_plan
from app.forecasting import calculate_coverage, forecast_message, projected_income_balances
from app import income_service
from app.income_detection_service import income_candidates
from app.income_reconciliation_service import resolve_income
from app.recurring_charge_service import decide_recurring_charge,recurring_charge_overview,upcoming_confirmed_charges
from app.models import (
    BillProfile,
    CategorizationRule,
    FinancialAccount,
    FinancialConnection,
    ImportRun,
    IncomeEvent,
    LiabilityStatement,
    PaymentMatch,
    ScheduledPayment,
    StatementUpload,
    Transaction,
    User,
)
from app.schemas import (
    AccountConnectionUpdate,
    AccountMergeRequest,
    BillCreate,
    BillCandidateIgnore,
    CategorizationRuleUpdate,
    LiabilityScheduleRequest,
    MonthlyPlanPaymentUpdate,
    LiabilityStatementCreate,
    ManualAccountCreate,
    IncomeCreate,
    PaymentCreate,
    PaymentUpdate,
    ReconcileRequest,
    RecurringChargeDecision,
    TransactionCategoryUpdate,
    TransactionMatchContextRequest,
    WhatIfRequest,
)

from app.login_routes import router as login_router
from app.simplefin_routes import router as simplefin_router
from app.simplefin_service import sync_health
from app.reconciliation_routes import router as reconciliation_router
from app.accounting_routes import router as accounting_router
from app.onboarding_routes import router as onboarding_router
from app.privacy_routes import router as privacy_router
from app.activity_routes import router as activity_router
from app.transaction_export_routes import router as transaction_export_router
from app.rate_limit import upload_rate_limit
from app.audit_service import record_audit

app = FastAPI(
    title="FinLeash API",
    version="1.0.0",
    docs_url=None if settings.environment == "production" else "/docs",
    redoc_url=None if settings.environment == "production" else "/redoc",
    openapi_url=None if settings.environment == "production" else "/openapi.json",
)
app.include_router(login_router)
app.include_router(simplefin_router)
app.include_router(reconciliation_router)
app.include_router(accounting_router)
app.include_router(onboarding_router)
app.include_router(privacy_router)
app.include_router(activity_router)
app.include_router(transaction_export_router)
app.add_middleware(
    CORSMiddleware,
    allow_origins=[settings.frontend_origin],
    allow_credentials=True,
    allow_methods=["GET", "POST", "PUT", "PATCH", "DELETE", "OPTIONS"],
    allow_headers=["Accept", "Content-Type"],
)
app.add_middleware(TrustedHostMiddleware, allowed_hosts=settings.trusted_host_list)
logger = logging.getLogger("finleash.api")


@app.middleware("http")
async def request_security(request: Request, call_next):
    if request.url.path.startswith("/api/v1") and request.method not in {
        "GET",
        "HEAD",
        "OPTIONS",
    }:
        origin = request.headers.get("origin")
        referer = request.headers.get("referer")
        supplied_origin = origin
        if not supplied_origin and referer:
            parsed = urlparse(referer)
            supplied_origin = f"{parsed.scheme}://{parsed.netloc}"
        must_supply = settings.is_secure_environment or COOKIE in request.cookies
        if (must_supply and not supplied_origin) or (
            supplied_origin and supplied_origin != settings.frontend_origin
        ):
            return JSONResponse(
                {
                    "error": {
                        "code": "CSRF_REJECTED",
                        "message": "Request origin is not allowed.",
                        "details": {},
                    }
                },
                status_code=403,
            )
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    if settings.is_secure_environment:
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    return response


@app.middleware("http")
async def observe_request(request: Request, call_next):
    supplied_request_id = request.headers.get("x-request-id", "")
    request_id = (
        supplied_request_id
        if re.fullmatch(r"[A-Za-z0-9._-]{8,64}", supplied_request_id)
        else str(uuid.uuid4())
    )
    token = request_id_context.set(request_id)
    started = time.monotonic()
    status_code = 500
    try:
        response = await call_next(request)
        status_code = response.status_code
        response.headers["X-Request-ID"] = request_id
        return response
    except Exception:
        logger.exception("Unhandled API request failure")
        raise
    finally:
        route_object = request.scope.get("route")
        route = getattr(route_object, "path", "unmatched")
        elapsed = time.monotonic() - started
        API_REQUESTS.labels(
            method=request.method, route=route, status=str(status_code)
        ).inc()
        API_LATENCY.labels(method=request.method, route=route).observe(elapsed)
        if status_code >= 500 and route != "unmatched":
            API_ERRORS.labels(method=request.method, route=route).inc()
        logger.info(
            "request_complete method=%s route=%s status=%s duration_ms=%.2f",
            request.method,
            route,
            status_code,
            elapsed * 1000,
        )
        request_id_context.reset(token)


@app.exception_handler(HTTPException)
async def error_handler(_, exc: HTTPException):
    from fastapi.responses import JSONResponse

    return JSONResponse(
        {"error": {"code": "REQUEST_FAILED", "message": str(exc.detail), "details": {}}},
        status_code=exc.status_code,
        headers=exc.headers,
    )


@app.get("/health")
def health():
    return {"status": "healthy"}


@app.get("/ready")
def ready(db: Session = Depends(get_db)):
    try:
        if not probe_dependencies(db):
            raise RuntimeError("dependency unavailable")
    except Exception:
        logger.warning("Readiness dependency probe failed")
        raise HTTPException(503, "Service is not ready.")
    return {"status": "ready"}


@app.get("/metrics", include_in_schema=False)
def metrics(request: Request, db: Session = Depends(get_db)):
    if not settings.metrics_enabled or not settings.metrics_auth_token:
        raise HTTPException(404, "Not found.")
    authorization = request.headers.get("authorization", "")
    expected = f"Bearer {settings.metrics_auth_token}"
    if not hmac.compare_digest(authorization, expected):
        raise HTTPException(
            401,
            "Metrics credentials are required.",
            headers={"WWW-Authenticate": "Bearer"},
        )
    try:
        refresh_health_metrics(db)
    except Exception:
        logger.warning("Health metric refresh failed")
    return Response(
        content=generate_latest(),
        headers={"Content-Type": CONTENT_TYPE_LATEST},
    )


def account_data(account: FinancialAccount, connection: FinancialConnection | None = None, balance_view: BalanceView | None = None) -> dict:
    biller, confidence = match_biller(account.name, account.institution_name or "")
    displayed_available = balance_view.amount if balance_view else account.available_balance
    displayed_balance = displayed_available if balance_view and account.kind in {"checking", "savings", "money_market"} else account.balance
    displayed_date = balance_view.as_of_date if balance_view else account.balance_as_of_date
    return {
        "id": account.id,
        "name": account.name,
        "kind": account.kind,
        "mask": account.mask,
        "balance": str(displayed_balance),
        "balance_as_of_date": displayed_date.isoformat()
        if displayed_date
        else None,
        "available_balance": str(displayed_available),
        "balance_source": balance_view.source if balance_view else account.data_source,
        "reported_balance": str(account.balance),
        "reported_available_balance": str(account.available_balance),
        "reported_balance_as_of_date": account.balance_as_of_date.isoformat() if account.balance_as_of_date else None,
        "balance_anchor_date": balance_view.anchor_date.isoformat() if balance_view and balance_view.anchor_date else None,
        "investment_balance": str(account.investment_balance),
        "original_balance": str(account.original_balance)
        if account.original_balance is not None
        else None,
        "reserve": str(account.reserve),
        "connection_mode": account.connection_mode,
        "data_source": account.data_source,
        "connection_id": account.connection_id,
        "connection_provider": connection.provider if connection else None,
        "connection_status": connection.status if connection else None,
        "sync_health": sync_health(connection) if connection else None,
        "last_sync_at": connection.last_sync_at.isoformat()
        if connection and connection.last_sync_at
        else None,
        "last_successful_sync_at": connection.last_successful_sync_at.isoformat()
        if connection and connection.last_successful_sync_at
        else None,
        "next_sync_at": connection.next_sync_at.isoformat()
        if connection and connection.next_sync_at
        else None,
        "institution_name": account.institution_name,
        "bank_login_url": account.bank_login_url,
        "customer_support": public_biller(biller, confidence) if biller else None,
    }


def dashboard(
    db: Session,
    hid: str,
    overrides: dict[str, tuple[date | None, Decimal | None]] | None = None,
    include_uncertain: bool = False,
) -> dict:
    forecast_date = date.today()
    period_start = forecast_date.replace(day=1)
    next_month = (
        date(forecast_date.year + 1, 1, 1)
        if forecast_date.month == 12
        else date(forecast_date.year, forecast_date.month + 1, 1)
    )
    period_end = next_month - timedelta(days=1)
    accounts = list(
        db.scalars(
            select(FinancialAccount).where(
                FinancialAccount.household_id == hid, FinancialAccount.is_active
            )
        )
    )
    payments = list(
        db.scalars(
            select(ScheduledPayment).where(
                ScheduledPayment.household_id == hid,
                ScheduledPayment.data_source != "seed",
                ScheduledPayment.earliest_withdrawal_date.between(period_start, period_end),
            )
        )
    )
    incomes = list(
        db.scalars(
            select(IncomeEvent).where(
                IncomeEvent.household_id == hid,
                IncomeEvent.data_source != "seed",
                IncomeEvent.expected_date.between(period_start, period_end),
            )
        )
    )
    transactions = list(
        db.scalars(
            select(Transaction).where(
                Transaction.household_id == hid,
                Transaction.direction == "credit",
                Transaction.pending.is_(False),
                Transaction.posted_date.between(period_start - timedelta(days=3), forecast_date),
            )
        )
    )
    account_map = {account.id: account for account in accounts}
    balance_views = {account.id: available_balance_view(db, hid, account) for account in accounts}
    resolved_incomes = [
        resolve_income(
            income, transactions,
            account_map[income.account_id].balance_as_of_date if income.account_id in account_map else None,
            forecast_date,
        )
        for income in incomes
    ]
    if overrides:
        for payment in payments:
            if payment.id in overrides:
                new_date, new_amount = overrides[payment.id]
                if new_date:
                    payment.earliest_withdrawal_date = new_date
                    payment.latest_withdrawal_date = new_date
                if new_amount:
                    payment.amount = new_amount
    coverages = []
    income_balance_map = {}
    for account in accounts:
        aps = [p for p in payments if p.account_id == account.id]
        ais = [i for i in resolved_incomes if i.account_id == account.id]
        coverages.extend(
            calculate_coverage(
                account, aps, ais, include_uncertain=include_uncertain,
                balance_as_of_date=balance_views[account.id].as_of_date,
                starting_balance=balance_views[account.id].amount,
            )
        )
        income_balance_map.update(
            projected_income_balances(
                account, aps, ais, include_uncertain, balance_views[account.id].as_of_date,
                balance_views[account.id].amount,
            )
        )
    cmap = {c.payment_id: c for c in coverages}
    ordered = []
    for p in sorted(payments, key=lambda x: x.earliest_withdrawal_date):
        c = cmap.get(p.id)
        ordered.append(
            {
                "id": p.id,
                "payee": p.payee,
                "amount": str(p.amount),
                "account_id": p.account_id,
                "scheduled_date": p.scheduled_date.isoformat(),
                "withdrawal_date": p.earliest_withdrawal_date.isoformat(),
                "latest_withdrawal_date": p.latest_withdrawal_date.isoformat(),
                "due_date": p.due_date.isoformat(),
                "method": p.method,
                "status": p.status,
                "confirmation_number": p.confirmation_number,
                "coverage_status": "paid" if p.status == "paid" else c.status if c else "funding_account_not_selected",
                "balance_after": str(c.balance_after) if c and p.status != "paid" else None,
                "reason": (f"Paid on {p.paid_date.isoformat()}." if p.paid_date else "Paid.")
                if p.status == "paid"
                else c.reason
                if c
                else "Choose a funding account before this can be evaluated.",
            }
        )
    counts = {key: 0 for key in ["covered", "tight", "at_risk", "unfunded"]}
    for row in ordered:
        if row["status"] == "paid":
            continue
        status = row["coverage_status"]
        counts[
            "covered"
            if status == "covered"
            else "tight"
            if status == "covered_but_tight"
            else "unfunded"
            if status in {"insufficient_funds", "funding_account_not_selected"}
            else "at_risk"
        ] += 1
    next_income = min(
        (
            i.expected_date
            for i in resolved_incomes
            if i.expected_date >= forecast_date
            and i.include_in_forecast
            and i.reliability in {"guaranteed", "high_confidence"}
        ),
        default=None,
    )
    active = [
        p for p in ordered if p["status"] not in {"paid", "cleared", "cancelled", "skipped", "failed"}
    ]
    total = sum(Decimal(p["amount"]) for p in active)
    lowest = min(
        (Decimal(p["balance_after"]) for p in ordered if p["balance_after"] is not None),
        default=sum(balance_views[a.id].amount for a in accounts),
    )
    status, reason = forecast_message(ordered)
    return {
        "period_start": period_start.isoformat(),
        "period_end": period_end.isoformat(),
        "status": status,
        "reason": reason,
        "available_cash": str(
            sum(balance_views[a.id].amount for a in accounts if a.kind in {"checking", "savings", "cash_management"})
        ),
        "scheduled_total": str(total),
        "next_income_date": next_income.isoformat() if next_income else None,
        "protected_reserve": str(sum(a.reserve for a in accounts)),
        "lowest_balance": str(lowest),
        "counts": counts,
        "accounts": [account_data(a, balance_view=balance_views[a.id]) for a in accounts],
        "payments": ordered,
        "subscriptions": upcoming_confirmed_charges(db,hid,period_start,period_end),
        "income": [
            {
                "id": i.id,
                "name": i.name,
                "date": i.expected_date.isoformat(),
                "amount": str(i.amount),
                "reliability": i.reliability,
                "account_id": i.account_id,
                "balance_after": str(income_balance_map[i.id]) if i.id in income_balance_map else None,
                "status": i.status,
                "actual_date": i.actual_date.isoformat() if i.actual_date else None,
                "reason": (
                    f"Matched to a posted deposit on {i.actual_date.isoformat()}."
                    if i.status == "received" and i.actual_date
                    else "Expected income; no matching posted deposit has been found yet."
                    if i.status == "not_confirmed"
                    else "No matching deposit was found within seven days of the expected date."
                    if i.status == "missing"
                    else "Estimated from prior deposits; the amount and date may change."
                ),
            }
            for i in sorted(resolved_incomes, key=lambda i: i.expected_date)
        ],
    }


@app.get("/api/v1/today")
def today(hid: str = Depends(household_id), db: Session = Depends(get_db)):
    return dashboard(db, hid)


def income_event_data(item: IncomeEvent) -> dict:
    return {
        "id": item.id,
        "name": item.name,
        "date": item.expected_date.isoformat(),
        "amount": str(item.amount),
        "reliability": item.reliability,
        "account_id": item.account_id,
        "status": item.status,
        "frequency": item.frequency,
        "series_id": item.series_id,
    }


@app.get("/api/v1/income")
def income_events(hid: str = Depends(household_id), db: Session = Depends(get_db)):
    items = db.scalars(
        select(IncomeEvent)
        .where(IncomeEvent.household_id == hid, IncomeEvent.data_source != "seed")
        .order_by(IncomeEvent.expected_date)
    )
    return [income_event_data(item) for item in items]


@app.get("/api/v1/income/candidates")
def detected_income_candidates(hid: str = Depends(household_id), db: Session = Depends(get_db)):
    return income_candidates(db, hid)


@app.get("/api/v1/bills/candidates")
def detected_bill_candidates(hid: str = Depends(household_id), db: Session = Depends(get_db)):
    return bill_candidates(db, hid)


@app.get("/api/v1/billers")
def biller_directory(q:str="",hid:str=Depends(household_id)):
    del hid
    return search_billers(q)


@app.get("/api/v1/bills")
def bills(hid: str = Depends(household_id), db: Session = Depends(get_db)):
    rows=db.scalars(select(BillProfile).where(BillProfile.household_id==hid,BillProfile.is_active).order_by(BillProfile.name))
    return [{"id":row.id,"name":row.name,"merchant_pattern":row.merchant_pattern,"category":row.category,"bill_type":row.bill_type,"amount_type":row.amount_type,"typical_amount":str(row.typical_amount),"frequency":row.frequency,"due_day":row.due_day,"default_account_id":row.default_account_id,"website_url":row.website_url,"biller_id":row.biller_directory_id,"payment_method":row.payment_method,"remaining_balance":str(row.remaining_balance) if row.remaining_balance is not None else None,"installments_remaining":row.installments_remaining} for row in rows]


@app.post("/api/v1/bills", status_code=201)
def add_bill(body: BillCreate, hid: str = Depends(household_id), db: Session = Depends(get_db)):
    try:
        bill=create_bill(db,hid,body)
    except ValueError as exc:
        raise HTTPException(422,str(exc)) from exc
    return {"id":bill.id,"name":bill.name}


@app.put("/api/v1/bills/{bill_id}")
def edit_bill(bill_id:str,body:BillCreate,hid:str=Depends(household_id),db:Session=Depends(get_db)):
    try:bill=update_bill(db,hid,bill_id,body)
    except ValueError as exc:raise HTTPException(422,str(exc)) from exc
    return {"id":bill.id,"name":bill.name}


@app.post("/api/v1/bills/candidates/ignore", status_code=204)
def ignore_bill(body:BillCandidateIgnore,hid:str=Depends(household_id),db:Session=Depends(get_db)):
    ignore_bill_candidate(db,hid,body.merchant_pattern)
    return Response(status_code=204)


@app.get("/api/v1/recurring-charges")
def recurring_charges(hid:str=Depends(household_id),db:Session=Depends(get_db)):
    return recurring_charge_overview(db,hid)


@app.post("/api/v1/recurring-charges/decision")
def review_recurring_charge(body:RecurringChargeDecision,hid:str=Depends(household_id),db:Session=Depends(get_db)):
    try:row=decide_recurring_charge(db,hid,body)
    except LookupError as exc:raise HTTPException(404,str(exc)) from exc
    except ValueError as exc:raise HTTPException(422,str(exc)) from exc
    return {"id":row.id,"status":row.status,"series_type":row.series_type}


@app.post("/api/v1/income", status_code=201)
def create_income(
    body: IncomeCreate, hid: str = Depends(household_id), db: Session = Depends(get_db)
):
    try:
        events = income_service.create_income(db, hid, body)
    except LookupError as exc:
        raise HTTPException(404, str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    return income_event_data(events[0]) | {"occurrences": len(events)}


@app.put("/api/v1/income/{income_id}")
def edit_income(
    income_id: str, body: IncomeCreate, hid: str = Depends(household_id), db: Session = Depends(get_db)
):
    try:
        events = income_service.update_income(db, hid, income_id, body, date.today())
    except LookupError as exc:
        raise HTTPException(404, str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    return income_event_data(events[0]) | {"occurrences": len(events)}


@app.delete("/api/v1/income/{income_id}", status_code=204)
def remove_income(
    income_id: str, scope: str = "one", hid: str = Depends(household_id), db: Session = Depends(get_db)
):
    if scope not in {"one", "all"}:
        raise HTTPException(422, "Scope must be one or all.")
    try:
        income_service.delete_income(db, hid, income_id, scope)
    except LookupError as exc:
        raise HTTPException(404, str(exc)) from exc
    return Response(status_code=204)


@app.get("/api/v1/accounts")
def accounts(hid: str = Depends(household_id), db: Session = Depends(get_db)):
    connections = {
        row.id: row
        for row in db.scalars(
            select(FinancialConnection).where(FinancialConnection.household_id == hid)
        )
    }
    rows = list(db.scalars(select(FinancialAccount).where(FinancialAccount.household_id == hid, FinancialAccount.is_active)))
    return [account_data(a, connections.get(a.connection_id), available_balance_view(db, hid, a)) for a in rows]


@app.post("/api/v1/accounts", status_code=201)
def create_account(
    body: ManualAccountCreate, hid: str = Depends(household_id), db: Session = Depends(get_db)
):
    return account_data(create_manual_account(db, hid, body))


@app.post(
    "/api/v1/accounts/statement-preview",
    dependencies=[Depends(upload_rate_limit)],
)
async def preview_statement_account(
    file: UploadFile = File(...), hid: str = Depends(household_id), db: Session = Depends(get_db)
):
    content = await file.read()
    if not content:
        raise HTTPException(422, "The uploaded file is empty.")
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(413, "Statement files must be 20 MB or smaller.")
    try:
        return preview_account_statement(db, hid, file.filename or "statement.pdf", content)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc


@app.get("/api/v1/institutions")
def institutions(q: str = "", hid: str = Depends(household_id)):
    del hid
    return search_institutions(q)


@app.post(
    "/api/v1/accounts/from-statement",
    status_code=201,
    dependencies=[Depends(upload_rate_limit)],
)
async def create_account_from_statement(
    name: str = Form(...),
    kind: str = Form(...),
    mask: str = Form(...),
    balance: Decimal = Form(...),
    available_balance: Decimal = Form(Decimal("0")),
    investment_balance: Decimal = Form(Decimal("0")),
    original_balance: Decimal | None = Form(None),
    reserve: Decimal = Form(Decimal("0")),
    institution_name: str = Form(""),
    bank_login_url: str = Form(""),
    institution_id: str = Form(""),
    existing_account_id: str = Form(""),
    file: UploadFile = File(...),
    hid: str = Depends(household_id),
    db: Session = Depends(get_db),
):
    allowed = {
        "checking",
        "savings",
        "credit_card",
        "money_market",
        "cash_management",
        "mortgage",
        "auto_loan",
        "buy_now_pay_later",
        "loan",
        "investment",
        "other",
    }
    if kind not in allowed:
        raise HTTPException(422, "Choose a supported account type.")
    if not re.fullmatch(r"[A-Za-z0-9]{2,4}", mask):
        raise HTTPException(422, "Enter the last 2–4 account characters.")
    trusted_institution = get_institution(institution_id) if institution_id else None
    if institution_id and not trusted_institution:
        raise HTTPException(422, "The selected institution is not in the FinLeash directory.")
    if trusted_institution:
        institution_name = trusted_institution.name
        bank_login_url = trusted_institution.login_url
    if bank_login_url:
        parsed_bank_url = urlparse(bank_login_url.strip())
        if (
            parsed_bank_url.scheme != "https"
            or not parsed_bank_url.netloc
            or parsed_bank_url.username
            or parsed_bank_url.password
        ):
            raise HTTPException(
                422, "Enter a public HTTPS bank URL without embedded credentials."
            )
    content = await file.read()
    if not content:
        raise HTTPException(422, "The uploaded file is empty.")
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(413, "Statement files must be 20 MB or smaller.")
    account = (
        db.scalar(
            select(FinancialAccount).where(
                FinancialAccount.id == existing_account_id,
                FinancialAccount.household_id == hid,
                FinancialAccount.is_active,
            )
        )
        if existing_account_id
        else None
    )
    if existing_account_id and not account:
        raise HTTPException(404, "The matching account was not found.")
    if account is None:
        account = FinancialAccount(
            household_id=hid,
            name=name.strip(),
            kind=kind,
            mask=mask,
            balance=abs(balance),
            balance_as_of_date=None,
            available_balance=abs(available_balance),
            investment_balance=abs(investment_balance),
            original_balance=abs(original_balance) if original_balance is not None else None,
            reserve=abs(reserve)
            if kind in {"checking", "savings", "money_market", "cash_management"}
            else Decimal("0"),
            connection_mode="manual",
            institution_name=institution_name.strip() or None,
            bank_login_url=bank_login_url.strip() or None,
            data_source="statement",
        )
        db.add(account)
        db.flush()
    elif trusted_institution:
        account.institution_name = trusted_institution.name
        account.bank_login_url = trusted_institution.login_url
    filename = file.filename or "statement.pdf"
    file_hash = hashlib.sha256(content).hexdigest()
    try:
        result = import_statement(db, hid, account, filename, content, balance)
        matched_bills = match_open_bill_payments(db, hid)
        result["bill_payments_matched"] = matched_bills
        db.add(
            StatementUpload(
                household_id=hid,
                account_id=account.id,
                filename=filename,
                file_format=result["detected_format"],
                file_hash=file_hash,
                file_size=len(content),
                status="completed",
                statement_end_date=date.fromisoformat(result["statement_end_date"])
                if result.get("statement_end_date")
                else None,
                transactions_added=result["added"],
                duplicates_skipped=result["duplicates"],
                data_source="statement",
            )
        )
        db.commit()
        db.refresh(account)
        return {
            "account": account_data(account),
            "import": result,
            "used_existing_account": bool(existing_account_id),
        }
    except ValueError as exc:
        db.rollback()
        raise HTTPException(422, str(exc)) from exc


@app.put("/api/v1/accounts/{account_id}")
def update_account(
    account_id: str,
    body: ManualAccountCreate,
    hid: str = Depends(household_id),
    db: Session = Depends(get_db),
):
    account = update_manual_account(db, hid, account_id, body)
    if not account:
        raise HTTPException(404, "Account not found.")
    return account_data(account)


@app.delete("/api/v1/accounts/{account_id}", status_code=204)
def delete_account(
    account_id: str, hid: str = Depends(household_id), db: Session = Depends(get_db)
):
    if not archive_account(db, hid, account_id):
        raise HTTPException(404, "Account not found.")
    return Response(status_code=204)


@app.post("/api/v1/accounts/merge")
def merge_account_records(
    body: AccountMergeRequest,
    user: User = Depends(current_user),
    hid: str = Depends(household_id),
    db: Session = Depends(get_db),
):
    try:
        return merge_accounts(db, hid, body.survivor_account_id, body.duplicate_account_id, user.id)
    except AccountMergeError as exc:
        raise HTTPException(422, str(exc)) from exc


@app.patch("/api/v1/accounts/{account_id}/connection")
def update_account_connection(
    account_id: str,
    body: AccountConnectionUpdate,
    hid: str = Depends(household_id),
    db: Session = Depends(get_db),
):
    account = update_manual_connection(db, hid, account_id, body)
    if not account:
        raise HTTPException(404, "Account not found.")
    return account_data(account)


@app.get("/api/v1/payments")
def payments(hid: str = Depends(household_id), db: Session = Depends(get_db)):
    return dashboard(db, hid)["payments"]


@app.post("/api/v1/payments", status_code=201)
def create_payment(
    body: PaymentCreate, hid: str = Depends(household_id), db: Session = Depends(get_db)
):
    if body.account_id and not db.scalar(
        select(FinancialAccount).where(
            FinancialAccount.id == body.account_id, FinancialAccount.household_id == hid
        )
    ):
        raise HTTPException(404, "Funding account not found.")
    p = ScheduledPayment(
        household_id=hid,
        account_id=body.account_id,
        payee=body.payee,
        amount=body.amount,
        minimum_amount=body.amount,
        scheduled_date=body.scheduled_date,
        earliest_withdrawal_date=body.expected_withdrawal_date,
        latest_withdrawal_date=body.expected_withdrawal_date,
        due_date=body.due_date,
        method=body.method,
        confirmation_number=body.confirmation_number,
        status="scheduled_at_biller" if body.confirmation_number else "planned",
    )
    db.add(p)
    db.commit()
    db.refresh(p)
    return {"id": p.id}


@app.patch("/api/v1/payments/{payment_id}")
def update_payment(
    payment_id: str,
    body: PaymentUpdate,
    hid: str = Depends(household_id),
    db: Session = Depends(get_db),
):
    p = db.scalar(
        select(ScheduledPayment).where(
            ScheduledPayment.id == payment_id, ScheduledPayment.household_id == hid
        )
    )
    if not p:
        raise HTTPException(404, "We could not find that scheduled payment.")
    for key, value in body.model_dump(exclude_unset=True).items():
        if key == "expected_withdrawal_date":
            p.earliest_withdrawal_date = value
            p.latest_withdrawal_date = value
        else:
            setattr(p, key, value)
    db.commit()
    return {"id": p.id, "status": p.status}


@app.get("/api/v1/payments/{payment_id}/match-candidates")
def match_candidates(payment_id: str, hid: str = Depends(household_id), db: Session = Depends(get_db)):
    if not db.scalar(select(ScheduledPayment.id).where(ScheduledPayment.id==payment_id,ScheduledPayment.household_id==hid)):
        raise HTTPException(404,"Payment not found.")
    return payment_match_candidates(db,hid,payment_id)

@app.post("/api/v1/transactions/payment-match-context")
def transaction_match_context(body:TransactionMatchContextRequest,hid:str=Depends(household_id),db:Session=Depends(get_db)):
    return transaction_payment_contexts(db,hid,body.transaction_ids)

@app.post("/api/v1/payments/{payment_id}/reconcile")
def reconcile(
    payment_id: str,
    body: ReconcileRequest,
    hid: str = Depends(household_id),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
):
    p = db.scalar(
        select(ScheduledPayment).where(
            ScheduledPayment.id == payment_id, ScheduledPayment.household_id == hid
        )
    )
    t = db.scalar(
        select(Transaction).where(
            Transaction.id == body.transaction_id, Transaction.household_id == hid
        )
    )
    if not p or not t:
        raise HTTPException(404, "Payment or bank transaction not found.")
    if p.account_id != t.account_id:
        if not body.confirm_account_change:
            raise HTTPException(409, "Confirm that this payment was made from a different account.")
        if t.id not in {candidate["id"] for candidate in payment_match_candidates(db,hid,p.id)}:
            raise HTTPException(409, "This transaction is not a plausible match for the payment.")
    if p.status in {"paid","cleared"} or db.scalar(select(PaymentMatch.id).where(PaymentMatch.payment_id==p.id)):
        raise HTTPException(409,"This payment is already matched.")
    if t.pending or t.direction!="debit" or t.is_transfer or db.scalar(select(PaymentMatch.id).where(PaymentMatch.transaction_id==t.id)):
        raise HTTPException(409,"This transaction is not eligible for matching.")
    db.add(
        PaymentMatch(
            household_id=hid,
            payment_id=p.id,
            transaction_id=t.id,
            match_type="exact" if abs(p.amount - t.amount) <= Decimal("0.01") else "suggested",
        )
    )
    p.status = "paid"
    p.paid_date = t.posted_date
    p.status_source = "matched_transaction"
    p.account_id = t.account_id
    record_audit(db,hid,user.id,"payment_matched",label=p.payee,date=t.posted_date.isoformat(),source="bank transaction")
    db.commit()
    return {"status": "paid", "paid_date": p.paid_date.isoformat(), "simulated": True}


@app.post("/api/v1/what-if")
def what_if(body: WhatIfRequest, hid: str = Depends(household_id), db: Session = Depends(get_db)):
    if not db.scalar(
        select(ScheduledPayment).where(
            ScheduledPayment.id == body.payment_id, ScheduledPayment.household_id == hid
        )
    ):
        raise HTTPException(404, "Payment not found.")
    return dashboard(
        db,
        hid,
        {body.payment_id: (body.expected_withdrawal_date, body.amount)},
        body.include_uncertain_income,
    )


@app.get("/api/v1/transactions")
def transactions(
    q: str = "",
    account_id: str = "",
    category: str = "",
    category_id: str = "",
    segment_id: str = "",
    amount: Decimal | None = None,
    selected_date: date | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    page: int = 1,
    page_size: int = 100,
    sort_by: str = "date",
    sort_dir: str = "desc",
    hid: str = Depends(household_id),
    db: Session = Depends(get_db),
):
    ensure_household_accounting(db, hid)
    db.commit()
    if page < 1:
        raise HTTPException(400, "Page must be 1 or greater.")
    if sort_by not in {"transaction", "date", "category", "amount"}:
        raise HTTPException(400, "Unsupported transaction sort field.")
    if sort_dir not in {"asc", "desc"}:
        raise HTTPException(400, "Sort direction must be asc or desc.")
    stmt = select(Transaction).where(Transaction.household_id == hid)
    if q:
        stmt = stmt.where(
            or_(
                Transaction.original_description.ilike(f"%{q}%"),
                Transaction.merchant.ilike(f"%{q}%"),
            )
        )
    if account_id:
        stmt = stmt.where(Transaction.account_id == account_id)
    if category:
        stmt = stmt.where(Transaction.category == category)
    if category_id:
        stmt = stmt.where(Transaction.category_id == category_id)
    if segment_id:
        stmt = stmt.where(Transaction.entity_id == segment_id)
    if amount is not None:
        stmt = stmt.where(Transaction.amount == abs(amount))
    if selected_date:
        stmt = stmt.where(Transaction.posted_date == selected_date)
    else:
        if date_from and date_to and date_from > date_to:
            raise HTTPException(400, "From date must be on or before To date.")
        if date_from:
            stmt = stmt.where(Transaction.posted_date >= date_from)
        if date_to:
            stmt = stmt.where(Transaction.posted_date <= date_to)
    page_size = min(max(page_size, 1), 100)
    total = db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    signed_amount = case(
        (Transaction.direction == "credit", Transaction.amount), else_=-Transaction.amount
    )
    sort_columns = {
        "transaction": func.lower(Transaction.merchant),
        "date": Transaction.posted_date,
        "category": func.lower(Transaction.category),
        "amount": signed_amount,
    }
    sort_column = sort_columns[sort_by]
    ordering = sort_column.asc() if sort_dir == "asc" else sort_column.desc()
    rows = list(
        db.scalars(
            stmt.order_by(ordering, Transaction.id.asc())
            .offset((page - 1) * page_size)
            .limit(page_size)
        )
    )
    return {
        "items": [
            {
                "id": t.id,
                "description": t.original_description,
                "merchant": t.merchant,
                "amount": str(t.amount),
                "direction": t.direction,
                "date": t.posted_date.isoformat(),
                "pending": t.pending,
                "category": t.category,
                "category_id": t.category_id,
                "segment_id": t.entity_id,
                "account_id": t.account_id,
                "data_source": t.data_source,
                "is_transfer": t.is_transfer,
            }
            for t in rows
        ],
        "page": page,
        "page_size": page_size,
        "total": total,
        "pages": max(1, (total + page_size - 1) // page_size),
    }


@app.patch("/api/v1/transactions/categories")
def update_transaction_categories(
    body: TransactionCategoryUpdate, hid: str = Depends(household_id), db: Session = Depends(get_db)
):
    try:
        ensure_household_accounting(db, hid)
        category_id = body.category_id
        if not category_id and body.category:
            categories = public_categories(db, hid)
            category_id = next((row["id"] for row in categories if row["name"] == body.category and row["allow_posting"]), None)
        if not category_id:
            raise AccountingError("Choose a transaction category.")
        return post_transaction_categories(db, hid, body.transaction_ids, category_id)
    except AccountingError as error:
        raise HTTPException(422, str(error)) from error


@app.get("/api/v1/categorization-rules")
def categorization_rules(hid: str = Depends(household_id), db: Session = Depends(get_db)):
    rows = list(
        db.scalars(
            select(CategorizationRule)
            .where(CategorizationRule.household_id == hid)
            .order_by(CategorizationRule.updated_at.desc())
        )
    )
    return [
        {
            "id": r.id,
            "pattern": r.pattern,
            "sample_description": r.sample_description,
            "category": r.category,
            "category_id": r.category_id,
            "is_active": r.is_active,
            "match_count": r.match_count,
            "last_matched_at": r.last_matched_at.isoformat() if r.last_matched_at else None,
            "created_at": r.created_at.isoformat(),
        }
        for r in rows
    ]


@app.patch("/api/v1/categorization-rules/{rule_id}")
def update_categorization_rule(
    rule_id: str,
    body: CategorizationRuleUpdate,
    hid: str = Depends(household_id),
    db: Session = Depends(get_db),
):
    rule = db.scalar(
        select(CategorizationRule).where(
            CategorizationRule.id == rule_id, CategorizationRule.household_id == hid
        )
    )
    if not rule:
        raise HTTPException(404, "Rule not found.")
    if body.category_id is not None:
        try:
            category = category_for(db, hid, body.category_id)
        except AccountingError as error:
            raise HTTPException(422, str(error)) from error
        rule.category_id = category.id
        rule.category = category.name
    elif body.category is not None:
        rule.category = body.category
    if body.is_active is not None:
        rule.is_active = body.is_active
    rule.dml_flag = "U"
    db.commit()
    return {"id": rule.id}


@app.delete("/api/v1/categorization-rules/{rule_id}", status_code=204)
def delete_categorization_rule(
    rule_id: str, hid: str = Depends(household_id), db: Session = Depends(get_db)
):
    rule = db.scalar(
        select(CategorizationRule).where(
            CategorizationRule.id == rule_id, CategorizationRule.household_id == hid
        )
    )
    if not rule:
        raise HTTPException(404, "Rule not found.")
    db.delete(rule)
    db.commit()
    return Response(status_code=204)


@app.post("/api/v1/statements/import", dependencies=[Depends(upload_rate_limit)])
async def upload_statement(
    account_id: str = Form(...),
    ending_balance: Decimal | None = Form(None),
    transaction_mode: str = Form("auto"),
    file: UploadFile = File(...),
    hid: str = Depends(household_id),
    db: Session = Depends(get_db),
):
    account = db.scalar(
        select(FinancialAccount).where(
            FinancialAccount.id == account_id,
            FinancialAccount.household_id == hid,
            FinancialAccount.is_active,
        )
    )
    if not account:
        raise HTTPException(404, "Account not found.")
    if transaction_mode not in {"auto", "review", "import_missing"}:
        raise HTTPException(422, "Choose review or import missing transactions.")
    online_connected = account.connection_id is not None
    review_only = online_connected and transaction_mode != "import_missing"
    content = await file.read()
    if not content:
        raise HTTPException(422, "The uploaded file is empty.")
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(413, "Statement files must be 20 MB or smaller.")
    filename = file.filename or "statement.csv"
    file_hash = hashlib.sha256(content).hexdigest()
    try:
        result = import_statement(db, hid, account, filename, content, ending_balance, import_transactions=not review_only)
        result["online_connected"] = online_connected
        matched_bills = match_open_bill_payments(db, hid) if not review_only else 0
        result["bill_payments_matched"] = matched_bills
        db.add(
            StatementUpload(
                household_id=hid,
                account_id=account.id,
                filename=filename,
                file_format=result["detected_format"],
                file_hash=file_hash,
                file_size=len(content),
                status="reviewed" if review_only else "completed",
                statement_end_date=date.fromisoformat(result["statement_end_date"])
                if result.get("statement_end_date")
                else None,
                transactions_added=result["added"],
                duplicates_skipped=result["duplicates"],
                data_source="statement",
            )
        )
        db.commit()
        return result
    except ValueError as exc:
        db.rollback()
        suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else "unknown"
        db.add(
            StatementUpload(
                household_id=hid,
                account_id=account.id,
                filename=filename,
                file_format=suffix[:12],
                file_hash=file_hash,
                file_size=len(content),
                status="failed",
                error_message=str(exc),
                data_source="statement",
            )
        )
        db.commit()
        raise HTTPException(422, str(exc)) from exc


@app.get("/api/v1/accounts/{account_id}/history")
def get_account_history(
    account_id: str,
    months: int = 12,
    hid: str = Depends(household_id),
    db: Session = Depends(get_db),
):
    history = account_history(db, hid, account_id, months)
    if history is None:
        raise HTTPException(404, "Account not found.")
    return history


@app.post(
    "/api/v1/liabilities/import",
    status_code=201,
    dependencies=[Depends(upload_rate_limit)],
)
async def import_liability_statement(
    account_id: str = Form(...),
    file: UploadFile = File(...),
    hid: str = Depends(household_id),
    db: Session = Depends(get_db),
):
    account = db.scalar(
        select(FinancialAccount).where(
            FinancialAccount.id == account_id,
            FinancialAccount.household_id == hid,
            FinancialAccount.is_active,
        )
    )
    if not account:
        raise HTTPException(404, "Liability account not found.")
    if not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(422, "Upload a PDF liability statement.")
    try:
        row = import_liability_pdf(db, hid, account, await file.read())
        return {"id": row.id}
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc


@app.get("/api/v1/liabilities")
def liabilities(hid: str = Depends(household_id), db: Session = Depends(get_db)):
    return latest_statements(db, hid)


@app.post("/api/v1/liabilities", status_code=201)
def add_liability(
    body: LiabilityStatementCreate, hid: str = Depends(household_id), db: Session = Depends(get_db)
):
    row = create_statement(db, hid, body)
    if not row:
        raise HTTPException(404, "Liability account not found.")
    return {"id": row.id}


@app.get("/api/v1/debt/strategy")
def debt_strategy(hid: str = Depends(household_id), db: Session = Depends(get_db)):
    return saved_strategy(db, hid)


@app.post("/api/v1/debt/strategy")
def create_debt_strategy(hid: str = Depends(household_id), db: Session = Depends(get_db)):
    return generate_strategy(db, hid)


@app.post("/api/v1/liabilities/{liability_id}/schedule", status_code=201)
def schedule_bill(
    liability_id: str,
    body: LiabilityScheduleRequest,
    hid: str = Depends(household_id),
    db: Session = Depends(get_db),
):
    payment = schedule_liability(db, hid, liability_id, body)
    if not payment:
        raise HTTPException(404, "Liability or checking account not found.")
    return {"id": payment.id, "status": payment.status}


@app.post("/api/v1/imports/preview", dependencies=[Depends(upload_rate_limit)])
async def preview_import(
    file: UploadFile = File(...), hid: str = Depends(household_id), db: Session = Depends(get_db)
):
    content = await file.read()
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = [
            [str(v) if v is not None else "" for v in row]
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 25), values_only=True)
        ]
        sheets.append({"name": ws.title, "rows": rows, "row_count": ws.max_row})
    key = hashlib.sha256(content).hexdigest()
    run = db.scalar(select(ImportRun).where(ImportRun.idempotency_key == key))
    if not run:
        run = ImportRun(
            household_id=hid,
            filename=file.filename or "workbook.xlsx",
            sheet_name=sheets[0]["name"],
            idempotency_key=key,
            report=f"{len(sheets)} sheets previewed",
        )
        db.add(run)
        db.commit()
    return {"import_id": run.id, "duplicate": run.status != "previewed", "sheets": sheets}


@app.put("/api/v1/monthly-plan/payment")
def save_monthly_plan_payment(
    body: MonthlyPlanPaymentUpdate, hid: str = Depends(household_id), db: Session = Depends(get_db)
):
    payment = (
        db.scalar(
            select(ScheduledPayment).where(
                ScheduledPayment.id == body.payment_id, ScheduledPayment.household_id == hid
            )
        )
        if body.payment_id
        else None
    )
    obligation = db.scalar(
        select(FinancialAccount).where(
            FinancialAccount.id == body.obligation_account_id,
            FinancialAccount.household_id == hid,
            FinancialAccount.is_active,
        )
    )
    checking = db.scalar(
        select(FinancialAccount).where(
            FinancialAccount.id == body.checking_account_id,
            FinancialAccount.household_id == hid,
            FinancialAccount.kind.in_({"checking", "cash_management"}),
            FinancialAccount.is_active,
        )
    )
    if not obligation and not payment:
        raise HTTPException(404, "Obligation account not found.")
    if not checking:
        raise HTTPException(422, "Choose an active checking or cash-management account as the funding source.")
    if body.payment_id and not payment:
        raise HTTPException(404, "Scheduled payment not found.")
    latest = (
        db.scalar(
            select(LiabilityStatement)
            .where(
                LiabilityStatement.household_id == hid,
                LiabilityStatement.account_id == obligation.id,
            )
            .order_by(LiabilityStatement.statement_date.desc())
        )
        if obligation
        else None
    )
    if payment is None:
        assert obligation is not None
        payment = ScheduledPayment(
            household_id=hid,
            account_id=checking.id,
            obligation_account_id=obligation.id,
            payee=obligation.name,
            amount=body.amount,
            minimum_amount=latest.minimum_payment if latest else body.amount,
            scheduled_date=body.scheduled_date,
            earliest_withdrawal_date=body.scheduled_date,
            latest_withdrawal_date=body.scheduled_date,
            due_date=body.due_date,
            method="external_bank_bill_pay",
            status=body.status,
            paid_date=body.paid_date,
            status_source="user",
            confirmation_number=body.confirmation_number,
            data_source="manual_plan",
        )
        db.add(payment)
    else:
        payment.account_id = checking.id
        payment.obligation_account_id = (
            obligation.id if obligation else payment.obligation_account_id
        )
        payment.amount = body.amount
        payment.scheduled_date = body.scheduled_date
        payment.earliest_withdrawal_date = body.scheduled_date
        payment.latest_withdrawal_date = body.scheduled_date
        payment.due_date = body.due_date
        payment.status = body.status
        payment.paid_date = body.paid_date if body.status == "paid" else None
        payment.status_source = "user"
        payment.confirmation_number = body.confirmation_number
        payment.dml_flag = "U"
    db.commit()
    db.refresh(payment)
    return {
        "id": payment.id,
        "status": payment.status,
        "paid_date": payment.paid_date.isoformat() if payment.paid_date else None,
    }


@app.get("/api/v1/monthly-plan")
def get_monthly_plan(
    year: int | None = None,
    month: int | None = None,
    hid: str = Depends(household_id),
    db: Session = Depends(get_db),
):
    today = date.today()
    try:
        return monthly_plan(db, hid, year or today.year, month or today.month)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc


@app.get("/api/v1/exports/monthly-plan.xlsx")
def export_plan(
    year: int | None = None,
    month: int | None = None,
    hid: str = Depends(household_id),
    db: Session = Depends(get_db),
):
    today = date.today()
    data = monthly_plan(db, hid, year or today.year, month or today.month)
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly Payment Plan"
    ws.append(
        [
            "Group",
            "Account",
            "Due",
            "Status",
            "Current balance",
            "Statement balance",
            "Statement date",
            "Minimum payment",
            "Scheduled date",
            "Payment amount",
            "Balance after payment",
        ]
    )
    for row in data["rows"]:
        ws.append(
            [
                row["group"],
                row["account_name"],
                row["due_date"],
                row["paid_status"],
                Decimal(row["current_balance"]) if row["current_balance"] is not None else None,
                Decimal(row["statement_balance"]) if row["statement_balance"] is not None else None,
                row["statement_date"],
                Decimal(row["minimum_payment"]),
                row["scheduled_date"],
                Decimal(row["payment_amount"]),
                Decimal(row["balance_after_payment"])
                if row["balance_after_payment"] is not None
                else None,
            ]
        )
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename=FinLeash-Monthly-Payment-Plan.xlsx"
        },
    )
